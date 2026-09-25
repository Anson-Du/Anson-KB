"""docx 解析器 - 使用 python-docx 读取 .docx 文件"""

import os
import re
from datetime import datetime
from typing import Optional

from .document import Document, Section


class DocxParser:
    """解析 .docx 文件，提取段落和表格文本"""

    def __init__(self, workspace_root: str):
        self.workspace_root = workspace_root

    def parse_file(self, file_path: str) -> Optional[Document]:
        """解析单个 .docx 文件"""
        full_path = os.path.join(self.workspace_root, file_path) if not os.path.isabs(file_path) else file_path

        # 跳过临时文件（以 ~$ 开头）
        if os.path.basename(file_path).startswith('~$'):
            return None

        if not os.path.exists(full_path):
            return None

        try:
            from docx import Document as DocxDocument
            docx = DocxDocument(full_path)
        except Exception:
            return None

        sections = []
        current_heading = "正文"
        current_lines = []

        # 提取段落
        for para in docx.paragraphs:
            text = para.text.strip()
            if not text:
                if current_lines:
                    sections.append(Section(heading=current_heading, content='\n'.join(current_lines)))
                current_lines = []
                continue

            # 检测标题
            if para.style and para.style.name and para.style.name.startswith('Heading'):
                if current_lines:
                    sections.append(Section(heading=current_heading, content='\n'.join(current_lines)))
                current_heading = text
                current_lines = []
            else:
                current_lines.append(text)

        # 处理最后一段
        if current_lines:
            sections.append(Section(heading=current_heading, content='\n'.join(current_lines)))

        # 提取表格
        for table in docx.tables:
            rows_text = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows_text.append(' | '.join(cells))
            if rows_text:
                sections.append(Section(
                    heading="表格",
                    content='\n'.join(rows_text),
                ))

        # 提取日期
        doc_date = self._extract_date(file_path, docx)
        title = self._extract_title(docx)

        doc = Document(
            path=file_path,
            date=doc_date,
            doc_type="docx",
            sections=sections,
            raw_length=sum(len(s.content) for s in sections),
            title=title,
        )
        return doc

    def _extract_date(self, file_path: str, docx) -> datetime:
        """从文件名或文档属性提取日期"""
        basename = os.path.basename(file_path)

        # 文件名日期模式
        date_patterns = [
            (r'(\d{4}-\d{2}-\d{2})', '%Y-%m-%d'),
            (r'(\d{4}\d{2}\d{2})', '%Y%m%d'),
            (r'(\d{2}-\d{2})', None),  # MM-DD，当年
        ]
        for pattern, fmt in date_patterns:
            match = re.search(pattern, basename)
            if match:
                date_str = match.group(1)
                try:
                    if fmt == '%Y-%m-%d':
                        return datetime.strptime(date_str, fmt)
                    elif fmt == '%Y%m%d':
                        return datetime.strptime(date_str, fmt)
                    else:
                        year = datetime.now().year
                        return datetime.strptime(f'{year}-{date_str}', '%Y-%m-%d')
                except ValueError:
                    pass

        # 文档属性
        try:
            props = docx.core_properties
            if props.modified:
                return props.modified
        except Exception:
            pass

        # 文件修改时间
        full_path = os.path.join(self.workspace_root, file_path) if not os.path.isabs(file_path) else file_path
        mtime = os.path.getmtime(full_path)
        return datetime.fromtimestamp(mtime)

    def _extract_title(self, docx) -> str:
        """从文档属性或首段提取标题"""
        try:
            title = docx.core_properties.title
            if title:
                return title
        except Exception:
            pass

        # 首段作为标题
        if docx.paragraphs:
            text = docx.paragraphs[0].text.strip()
            if text:
                return text

        return ""
