# -*- coding: utf-8 -*-
# 生成脚本：AI产业生态合作与深圳市场拓展策略讨论（清理版）.docx
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from data1_0829 import SUMMARY_GROUPS, CHAPTERS_P1
from data2_0829 import CHAPTERS_P2, TRANSCRIPT_NOTES

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r'D:\Datawhale\AI产业生态合作与深圳市场拓展策略讨论（清理版）.docx'

SPEAKER_COLORS = {
    'Datawhale创始人': RGBColor(0xC0, 0x00, 0x00),  # 外部对话方 深红
    'Anson': RGBColor(0x00, 0x66, 0xB0),             # 本人 蓝
}

doc = Document()

# 页面设置
sec = doc.sections[0]
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.18)
sec.right_margin = Cm(3.18)

# 正文样式
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
style.paragraph_format.line_spacing = 1.5


def set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_after = Pt(6)


def add_meta(lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_font(r, size=10, color=RGBColor(0x80, 0x80, 0x80), italic=True)


def add_section_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=14, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    # 底部下划线分隔
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pPr.makeelement(qn('w:bottom'), {})
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_chapter(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=13, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.makeelement(qn('w:pBdr'), {})
    bottom = pPr.makeelement(qn('w:bottom'), {})
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_speaker_line(speaker, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(speaker + '：')
    set_font(r1, size=11, bold=True, color=SPEAKER_COLORS.get(speaker, RGBColor(0, 0, 0)))
    r2 = p.add_run(text)
    set_font(r2, size=11)


def add_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=10, color=RGBColor(0x80, 0x80, 0x80))


def add_summary_group(group_title, points):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(group_title)
    set_font(r, size=12, bold=True, color=RGBColor(0x40, 0x40, 0x40))
    for i, pt in enumerate(points, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run('%d. %s' % (i, pt))
        set_font(r, size=11)


# ============ 文档内容 ============
add_title('AI产业生态合作与深圳市场拓展策略讨论')
add_meta([
    '时间：2026年8月下旬（按转写语境推断）',
    '时长：约1小时16分钟',
    '形式：一对一深度交流',
    '参会人：Datawhale创始人、Anson',
])

add_section_heading('各方观点与结论总结')
for g_title, points in SUMMARY_GROUPS:
    add_summary_group(g_title, points)

add_section_heading('对话正文')
add_note('以下为清理后的对话记录，按议题分章。原始转写质量较差，已做语义重构；发言人均按内容语境归因。')

for title, turns in CHAPTERS_P1 + CHAPTERS_P2:
    add_chapter(title)
    for sp, tx in turns:
        add_speaker_line(sp, tx)

add_section_heading('转写处理说明')
for note in TRANSCRIPT_NOTES:
    add_note('· ' + note)

doc.save(OUT)
print('saved:', OUT)
