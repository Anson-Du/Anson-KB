# -*- coding: utf-8 -*-
"""
08-17 余杭和杭州政策沟通 - 会议纪要 Word 文档生成脚本
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ========== 配置区 ==========

output_path = r"D:\RoboX\03_政府合作\余杭\2026-08-17 16.33 余杭和杭州政策沟通（会议纪要）.docx"

# 发言人颜色映射
SPEAKER_COLORS = {
    '桂林': RGBColor(0xC0, 0x00, 0x00),    # 余杭区政策对接人 - 深红
    'Anson': RGBColor(0x70, 0x30, 0xA0),   # Anson - 紫色
    '碧莹': RGBColor(0x00, 0x60, 0xA0),    # 碧莹 - 深蓝
    '花花': RGBColor(0x00, 0x80, 0x80),    # 花花 - 青色
}

# ========== 格式引擎 ==========

COLOR_HEADING = RGBColor(0x1F, 0x4E, 0x79)
COLOR_NORMAL = RGBColor(0x00, 0x00, 0x00)
COLOR_NOTE = RGBColor(0x80, 0x80, 0x80)

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
pf = style.paragraph_format
pf.space_after = Pt(6)
pf.line_spacing = 1.5


def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.space_before = Pt(6)


def add_heading(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_dialogue(speaker, text):
    p = doc.add_paragraph()
    run_s = p.add_run(f'【{speaker}】')
    run_s.font.bold = True
    run_s.font.size = Pt(11)
    run_s.font.color.rgb = SPEAKER_COLORS.get(speaker, COLOR_NORMAL)
    run_s.font.name = '微软雅黑'
    run_s.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run_t = p.add_run(text)
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = COLOR_NORMAL
    run_t.font.name = '微软雅黑'
    run_t.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5


def add_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = COLOR_NOTE
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(1)


# ========== 文档内容 ==========

add_title('余杭区和杭州市AI企业政策申报与算力补贴沟通纪要')

add_note('参会人：桂林、Anson、碧莹、花花')
add_note('时间：2026年8月17日')

# ============================
add_heading('一、算力券政策详解（余杭区 vs 杭州市级）')
# ============================

add_dialogue('桂林', '今年6月到7月是上一轮周期，今年正式开始后会改成先申请再兑付。余杭区今年开放算力券，总额5000万，企业能拿到的算力补贴是算力采购金额的50%，单企业补贴上限国产算力200万、非国产150万。但有限制：必须采购余杭区算力伙伴库内的资源，目前入库的算力服务商有几家。申报通道在"余审新"平台上，时间分两批：今年10月一批，明年4月一批。')

add_dialogue('桂林', '杭州市也有算力补贴，总额比余杭区高，单企业上限800万，但补贴比例低：国产30%、非国产20%。关键是市区两级"就高不重复"——同一笔算力采购支出，在区级和市级之间只能选高的一档，同一份合同不能在两边申请。')

add_dialogue('Anson', '那就是两边都可以申请，就高不重复，一笔合同只能用来申请一个地方的算力补贴。每个区市两级各有比例和上限要求，所以合同需要设计——哪些合同拿到市里申请，哪些拿在区里。')

add_dialogue('桂林', '对，同一份合同不能在两边申请。如果算力使用量比较大，建议拆成两个合同。但有个顺序问题：如果先申请余杭区，再申请杭州市的，不影响；如果先申请杭州市的，再申请余杭区时可能已经超了区级上限。建议按50%来拆，超出的部分放到另一份合同，两个都可以尝试申请。')

add_dialogue('Anson', '只有裸算力才符合要求吗？推理也行，但容器之类的就不行了？')

add_dialogue('桂林', '对，智算云服务器、智能云服务这些可以。存储、网络安全类的不行，限制比较多。')

# ============================
add_heading('二、算力采购合同注意事项')
# ============================

add_dialogue('Anson', '英伟达的卡现在还要求合同里写什么卡吗？阿里云不敢在合同里写英伟达的卡。之前市经信局要求合同必须写清楚谁的卡。')

add_dialogue('桂林', '这个可能需要算力服务商协助提供相关材料。国内卡型补贴额度高一点，国外的低一些。如果总额不大，余杭区划算；如果使用量大，市级的上限更高。')

add_dialogue('Anson', '这次先内部交流，把问题收敛完之后，再跟阿里云的人一起商量。')

add_dialogue('桂林', '包括时间线之前跟阿里云也聊过，要拿补贴必须配合提供相关材料。我们也可以帮你们去跟阿里云沟通。')

add_dialogue('Anson', '阿里云的注册地不在余杭吧？')

add_dialogue('桂林', '对，阿里云注册地在西湖区，所以合同不能直接跟阿里云签，要跟余杭区的代理商签——余杭区大数据经营有限公司。余杭区政策既要扶持区内企业，也要扶持区内算力服务商。跟代理商签合同时，要说明需要申请算力券，他们后期要协助提供相关材料。')

add_dialogue('Anson', '阿里如果投我们，会层层穿透去看吗？')

add_dialogue('桂林', '间接的会看，但如果穿透下来占股比例不大、不是实控人，一般还好。直接或间接同为第三方控制、实控人是同一家就不行。')

# ============================
add_heading('三、算力合同细节与发票要求')
# ============================

add_dialogue('Anson', '接下来陆续要签合同了，需要对齐一下细节：合同到底怎么写，政府要什么材料，能不能做到。之前主要卡点有几个：第一，政府要求合同写明什么卡，阿里云不敢写；第二，合同周期和付款周期必须映射到同一阶段，流水和付款都要在补贴周期内。')

add_dialogue('桂林', '必须在合同有效期内，补贴的是年度内实际使用的部分。如果合同是之前签的、履约周期很长，只要包含使用周期就行。')

add_dialogue('Anson', '票据呢？收款方提供的票据有没有具体时间要求？')

add_dialogue('桂林', '发票需要备注写明合同号和对应的时间段，要很清楚。')

add_dialogue('Anson', '这些细节特别重要，比如发票开晚一点开到2027年去了，那2026年的补贴就拿不了。我们得跟阿里云配合好，流水、发票、合同都要符合要求。')

# ============================
add_heading('四、模型券与模型备案补贴')
# ============================

add_dialogue('Anson', '这个公司全流程是具身智能的，从模型到算力到数据到产品都有。每一个跟人工智能有关的政策都有关系。公司有自主研发的模型，不止一个——包括小脑模型、VLA (Vision-Language-Action，视觉-语言-动作) 模型、大脑模型。消耗算力就是为了调试这些模型。')

add_dialogue('桂林', '目前有两项可以申请：一是模型券，二是应用示范项目。模型备案的补贴要等9月份左右看有没有新通知。模型券按模型服务采购金额的20%补助，如果采购的是余杭区模型服务商的服务，比例可以到30%，上限100万，仅限API (Application Programming Interface，应用程序接口) 订阅调用。')

add_dialogue('Anson', '模型备案区里和市里是不是都有奖励？')

add_dialogue('桂林', '区级垂类模型应用示范项目最高300万，大走廊的模型备案奖励资金还没确认，之前应该是100万左右。')

add_dialogue('Anson', '先告诉我们备案能拿到多少钱，我们才有动力去备案。你帮我们盯着模型备案的补贴通知，园区目前应该还没有公司做过备案。')

add_dialogue('桂林', '垂类模型应用示范项目要求：开发的垂类大模型必须在余杭区有落地应用——跟余杭区企业签订落地合同、实际产生应用。年度选5个，不超过5个的都能拿，超过的评审。评审通过可拿研发投入的30%。材料里要求提供模型备案证明。')

# ============================
add_heading('五、资源券')
# ============================

add_dialogue('桂林', '资源券是今年新出的，额度较小，上限5万到10万。每月第一个工作日线上抢。但有个限制：有白名单，目前200多家企业，基本上是规上国高企业。作为模型券的补充政策。额度少，且只针对国内模型，得用指定平台采购。现阶段可能不太适用。')

# ============================
add_heading('六、领军人才与双创评审政策')
# ============================

add_dialogue('桂林', '评审类政策是比较稳的。一个是领军人才项目，一个是双创评审项目。领军人才：三年最高1200万（单年400万），按实际研发投入30%或设备采购20%补贴，房租补80%。双创评审：三年最高600万，研发投入30%，房租60%。双创对项目本身有要求，但对申请人要求低一些——持股15%以上就行。领军对申请人要求高：必须博士以上学位、第一大股东或最大自然人持股15%以上。')

add_dialogue('Anson', '领军人才对年龄有要求吗？')

add_dialogue('桂林', '年龄没要求，但学位有要求，一般博士以上。优秀人才也要求博士以上，但放宽的可能性更大。很多人卡在人的资格条件上达不到。')

add_dialogue('碧莹', '主要是博士学历这个要求。')

add_dialogue('Anson', '联创的股份如果不到15%的话比较麻烦。这些条件里哪些是可以松动的？还是说钢板一块？')

add_dialogue('桂林', '学历放松到硕士可以跟政府谈，但自然去报比较难。如果是硕士、项目特别优秀、融资比较大，是有先例可以谈的。本科的话就不太可能了。但归根结底，政府要合规，除非主要领导打招呼，否则担心回头审计。现在国家发了负面清单，整体政策在收紧。')

add_dialogue('Anson', 'C类为什么难？博士才是D类是吧？')

add_dialogue('桂林', '对，双创的话补贴是领军的一半，但要求也低。最开始余杭区没有双创项目，因为领军对人才限制太死，很多优秀企业拿不到，才衍生出双创。')

add_dialogue('Anson', '帮我们盯着领军的节奏，及时提醒。我们肯定要申请人才系列的。另外把领军和优秀人才的条件告诉碧莹，我们得组织资料看看能满足哪个，硬条件是什么。')

add_dialogue('桂林', '好的。领军的话业务板块可以满足具身智能，时间要求是落地余杭两年以内。下一批双创评审预计11月份，领军的话下一批可能也在那个时间。你们的项目报双创的话，通过概率90%以上。')

# ============================
add_heading('七、一事一议政策现状（全国统一大市场背景）')
# ============================

add_dialogue('Anson', '今年年初财政部出了一个文，执行统一大市场政策，不允许地方政府有一事一议政策，原来有的全得停。如果不想停，第一要主动交代原有的一事一议事项，第二要说明为什么还要一事一议。现在地方政府不会玩了——原来的政策不让用，新的政策也没有，都在观望深圳。深圳有立法权，每次政策调整都是深圳先行先试，全国效仿。今年就是青黄不接、观望阶段。')

add_dialogue('Anson', '甚至有地方钱都发完了，书记带着区长去求企业把钱吐出来。当官的在乎乌纱帽，今年一事一议难度非常大，还不如明年谈。今年是国家严打阶段。')

add_dialogue('桂林', '评审类政策完全没问题，通过评审不怕回头看。但一事一议的政策，包括已经迁出去还没执行的，大概率都要取消。')

add_dialogue('Anson', '碧莹你能理解了吧？一事一议今年难度大。')

add_dialogue('碧莹', '嗯。')

# ============================
add_heading('八、公司当前状态与申报节奏')
# ============================

add_dialogue('桂林', '你们现在在余杭有拿过其他补贴吗？')

add_dialogue('碧莹', '什么都没拿，公司刚成立。')

add_dialogue('桂林', '第三批双创评审今天刚结束，下一批预计10月底11月。企业注册时间得是落地余杭两年以内。你们如果现在或明年来考虑也刚好。')

add_dialogue('碧莹', '我们现在人员还没几个，核心就几个。股权变更、投资都还在谈，可能要到9月。')

add_dialogue('桂林', '补贴是先交后返，根据企业实际投入按比例返还。研发投入、租金补贴是创新企业早期比较大头的。如果业务稍微往前推一推再来考虑也可以。')

add_dialogue('Anson', '我们不是不为业务打基础，是人员还没进来。到9月以后肯定要做中创，对人员有要求。')

# ============================
add_heading('九、后续安排')
# ============================

add_dialogue('Anson', '今年先申请什么？算力券、模型备案的补贴盯紧。接下来陆续要签算力合同，合同细节、发票要求都得跟阿里云对齐。你这边回去想一下，整个相关补贴项还有哪些可以给我们建议。')

add_dialogue('桂林', '好的，我会把跟你们可以匹配的政策整理一下发过来，包括领军人才的具体条件。新的政策出来也会及时通知。')

add_dialogue('Anson', '另外一件事：今天到了一批机械臂，打算放到外面工位地胶板位置，因为会议室有静电。需要安排安全巡检，看看能不能装个摄像头。我先带你们去看看场地。')

# 保存
doc.save(output_path)
print(f"已生成: {output_path}")
