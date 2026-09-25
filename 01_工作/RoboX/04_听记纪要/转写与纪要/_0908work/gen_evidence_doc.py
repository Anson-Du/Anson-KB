# -*- coding: utf-8 -*-
"""生成《09-08 长虹洽谈 发言人归因证据表》Word 文档"""
import importlib.util, os, re
from collections import defaultdict, Counter
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = 'D:/RoboX/04_听记纪要/转写与纪要/'
DBLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xB2, 0x22, 0x22)

# --- 复用 audit 逻辑：以子进程运行 audit_evidence2.py 并捕获输出 ---
import subprocess
PY = 'C:/Users/Lenovo/.workbuddy/binaries/python/envs/default/Scripts/python.exe'
report = subprocess.run([PY, BASE + '_0908work/audit_evidence2.py'],
                        capture_output=True, text=True, encoding='utf-8').stdout

# 从 report 中解析矩阵
matrix_rows = []
for ln in report.split('\n'):
    if ln.startswith('田明') or ln.startswith('Anson') or ln.startswith('邹宇') or ln.startswith('曾昊山') or ln.startswith('周总') or ln.startswith('童哥') or ln.startswith('张维') or ln.startswith('万涛'):
        matrix_rows.append(ln)

def setf(run, size=11, bold=False, color=DARK, name='微软雅黑'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = color

def para(doc, text, size=11, bold=False, color=DARK, indent=0, align=WD_ALIGN_PARAGRAPH.LEFT, space=4):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(space)
    if indent: p.paragraph_format.left_indent = Cm(indent)
    p.alignment = align
    setf(p.add_run(text), size, bold, color)
    return p

def h1(doc, text, size=15, color=DBLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); setf(r, size, True, color); r.font.underline = True
    return p

def shade(cell, hexc):
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), hexc)
    cell._tc.get_or_add_tcPr().append(sh)

def borders(table):
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tb = OxmlElement('w:tcBorders')
            for b in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement(f'w:{b}')
                e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'BFBFBF')
                tb.append(e)
            tcPr.append(tb)

doc = Document()
st = doc.styles['Normal']
st.font.name = '微软雅黑'; st._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑'); st.font.size = Pt(11)
sec = doc.sections[0]
for a in ['left_margin', 'right_margin', 'top_margin', 'bottom_margin']:
    setattr(sec, a, Cm(2.3))

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
setf(p.add_run('09-08 睦灵科技 × 长虹集团洽谈'), 18, True, DBLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
setf(p.add_run('转写发言人归因证据表（用于人工裁决）'), 12, True, GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
setf(p.add_run('2026-09-20 · 依据源转写 1,063 个发言人块 · k-gram 投票法'), 9, False, GRAY)

h1(doc, '一、方法与判据')
para(doc, '以源转写的全部发言人块建立 6-gram 倒排索引，把清理版每一条问答（QA）回投到源文，'
          '按【源发言人编号】聚合投票，取得票最高者为建议归属。', 10.5)
para(doc, '指标：share = 该编号得票 / 全部命中票；lcs = QA 与该编号语料的最长公共子串长度。', 10.5)
para(doc, '裁决线：share ≥ 0.50 且 lcs ≥ 12 → 高置信；share ≥ 0.35 → 中置信；其余作低置信（不作数）。', 10.5)
para(doc, '样本量：清理版 QA 共 196 条，其中可定位 78 条、完全无命中 118 条、低置信 1 条。', 10.5, True, RED)

h1(doc, '二、证据矩阵（清理版 speaker × 源发言人编号）')
headers = ['清理版 speaker', '1', '2', '3', '4', '5', '6', '7', '合计']
sub = ['—', '曾昊山', '田明', '邹宇Tony', '待定', '周怡成', 'Anson', '童哥', '']
t = doc.add_table(rows=1 + len(matrix_rows) + 1, cols=len(headers))
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(headers):
    c = t.rows[0].cells[i]
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    setf(c.paragraphs[0].add_run(h), 9.5, True, RGBColor(0xFF, 0xFF, 0xFF))
    shade(c, '1F4E79')
for i, s in enumerate(sub):
    c = t.rows[1].cells[i]
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    setf(c.paragraphs[0].add_run(s), 9, False, GRAY)
    shade(c, 'DCE6F1')
for r, line in enumerate(matrix_rows, start=2):
    cells = re.split(r'\s{2,}', line.strip())
    for i in range(len(headers)):
        c = t.rows[r].cells[i]
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
        setf(c.paragraphs[0].add_run(cells[i] if i < len(cells) else ''), 9.5, i == 0)
borders(t)
para(doc, '注：括号内为该格"高置信"条数；· 表示无匹配。', 9, False, GRAY, space=8)

h1(doc, '三、裁决结论')
rows = [
    ('发言人 1', '曾昊山（总助）', '✅ 确认', '11 条命中、其中 7 条高置信；内容为介绍贾总/周总/童哥等长虹同事，符合主持人身份'),
    ('发言人 2', '田明', '✅ 确认', '源文自报"我主要就是产品的负责人田明"；10 条命中、3 条高置信'),
    ('发言人 3', '邹宇（Tony／饺子）', '✅ 确认', '内容为协议条款（附件四、5.4/5.5 验收、竞品、MOU 定稿、5 份协议），系业务方案与合同对接人'),
    ('发言人 4', '长虹方（姓名待确认）', '⚠ 存疑', '内容显示为长虹侧主导人（为张维辩护、牵头请财务与合规、主导宇树/乐聚本体采购、提及"我们的 ODC"），'
                                        '且其称"那贾总…"表明本人不是贾澜鹏'),
    ('发言人 5', '周总（周怡成）', '✅ 确认', '源文表头标注；6 条命中、4 条高置信；内容为财务视角尽调（现金流/注册资本/折旧/战略价值）'),
    ('发言人 6', 'Anson', '✅ 确认', '源文表头标注；对接收尾 action（"跟曾总对齐"）符合产业生态与 GA 角色'),
    ('发言人 7', '童哥（合规／法务）', '✅ 确认', '锚点"公平竞争法里面的规则"落在发言人 7；4 条高置信；内容为合规/内部研发流程'),
]
t2 = doc.add_table(rows=1 + len(rows), cols=4)
for i, h in enumerate(['源槽位', '裁决归属', '状态', '证据']):
    c = t2.rows[0].cells[i]
    setf(c.paragraphs[0].add_run(h), 9.5, True, RGBColor(0xFF, 0xFF, 0xFF)); shade(c, '1F4E79')
for r, row in enumerate(rows, 1):
    for i, v in enumerate(row):
        c = t2.rows[r].cells[i]
        setf(c.paragraphs[0].add_run(v), 9, i == 1, RED if (i == 2 and '⚠' in v) else DARK)
borders(t2)

h1(doc, '四、待人工裁决事项')
para(doc, '1. 发言人 4 的真实姓名。证据指向"长虹侧统筹人"（后被辞作贾总的对接人），'
          '但清理版目前以"贾澜鹏（贾总，运营总监）"入文——两者是否同一人，需与会人确认。', 10.5, color=RED)
para(doc, '2. "张总（财务合规）"这一早期假设在证据中未获支持：发言人 7 的高置信归属为童哥。'
          '源文中田明寒暄确有"周总，张总，王总"的称呼，张总应另有其人（可能为张维的尊称或未发言者）。', 10.5, color=RED)
para(doc, '3. 会议地点：清理版记"杭州"（依据田明开场"辛苦大家远道而来"），'
          '但发言人 4 有"因为您刚过来吧"等表述，建议一并复核。', 10.5, color=RED)

h1(doc, '五、质量提示：正文前 5 章为改写而非逐句清理')
para(doc, '经词频核验，清理版正文前 5 章存在较多改写成分——"三方面""融资节奏""首笔""投后""蚂蚁灵波"'
          '"数字水印""入表""优先供给""独家供货""路与车""框架协议""子合同""战略价值报告"等表述'
          '在源转写中零命中；第 6–14 章（协议逐条对接）则与源文高度贴合。', 10.5)
para(doc, '影响：前 5 章可作为"议题纪要"阅读，但不能作为发言原话引用；若需对外提交或作为洽谈留痕，'
          '建议按源文重建为逐句清理版。', 10.5, True, RED)

out = BASE + '09-08 长虹洽谈_发言人归因证据表.docx'
doc.save(out)
print('✅', out, f'{os.path.getsize(out)/1024:.1f} KB, 段落 {len(doc.paragraphs)}')
