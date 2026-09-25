# -*- coding: utf-8 -*-
"""0827深创投江思贤交流清理版 - 主生成脚本"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from data1_0827 import SUMMARY_GROUPS, CHAPTERS_P1
from data2_0827 import CHAPTERS_P2

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   '0827 深创投江思贤交流（清理版）.docx')

SPEAKER_COLORS = {
    '江思贤': RGBColor(0xC0, 0x00, 0x00),   # 外部嘉宾 深红
    '明修':   RGBColor(0x70, 0x30, 0xA0),   # 创始人/CTO 紫
    '田明':   RGBColor(0x00, 0x80, 0x80),   # 联合创始人 青
    '信庭':   RGBColor(0xC5, 0x5A, 0x11),   # 团队成员 橙棕
    'Anson':  RGBColor(0x00, 0x66, 0xB0),   # 蓝色
}

doc = Document()
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

for sec in doc.sections:
    sec.top_margin = Cm(2.54); sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18); sec.right_margin = Cm(3.18)

def _set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    if color is not None: run.font.color.rgb = color

def add_title(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(text), size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_after = Pt(6)

def add_meta(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    _set_font(p.add_run(text), size=10, italic=True, color=RGBColor(0x80, 0x80, 0x80))

def add_part_heading(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14)
    _set_font(p.add_run(text), size=13, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

def add_chapter(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    run = p.add_run(text)
    _set_font(run, size=12.5, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    p.paragraph_format.border_bottom = True if False else None
    # 底部下划线分隔（用段落边框实现）
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom); pPr.append(pBdr)

def add_speaker_line(sp, tx):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(sp + '：')
    _set_font(run, bold=True, color=SPEAKER_COLORS.get(sp, RGBColor(0, 0, 0)))
    run2 = p.add_run(tx)
    _set_font(run2)

def add_note(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    _set_font(p.add_run(text), size=10, italic=True, color=RGBColor(0x80, 0x80, 0x80))

def add_summary_point(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    _set_font(p.add_run('• ' + text))

# ================= 文档组装 =================
add_title('具身智能创业团队技术路线与融资策略交流（清理版）')
add_meta('时间：2026年8月27日（转写时长约1小时25分钟）')
add_meta('地点：线下交流（杭州）')
add_meta('参会人：明修、田明、信庭、Anson、江思贤')
add_meta('说明：本文档由录音转写清理生成，口头语与闲聊已删除，转写错误已修正；' 
         '原始转写质量较差，部分内容基于上下文语义重构，个别细节（数字/人名/机构）如与实际有出入，请以录音为准。')

add_part_heading('沟通总结')
for gtitle, points in SUMMARY_GROUPS:
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    _set_font(p.add_run(gtitle), size=11.5, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    for pt in points:
        add_summary_point(pt)

doc.add_page_break()

add_part_heading('对话正文（按议题整理）')
for title, lines in CHAPTERS_P1:
    add_chapter(title)
    for sp, tx in lines:
        add_speaker_line(sp, tx)

add_part_heading('对话正文·续')
for title, lines in CHAPTERS_P2:
    add_chapter(title)
    for sp, tx in lines:
        add_speaker_line(sp, tx)

add_note('注：转写中"蚂蚁宁波/灵波"均指蚂蚁灵波；"天机/天玑5G臂"指天玑机械臂与舞肌灵巧手；'
         '"语速/榆树"指宇树机器人；"阿里郎院/安利院"指阿里达摩院。'
         '江思贤所述海康机器人市值等数字为口述口径，未经核实。')

doc.save(OUT)
print('saved:', OUT)
