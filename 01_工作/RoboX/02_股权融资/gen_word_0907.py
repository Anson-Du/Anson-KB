# -*- coding: utf-8 -*-
"""
0907深创投投委方爱华路演清理版 - 生成脚本
从 data1/data2/data3 读取数据，生成清理版Word文档
"""
import sys
import os

# 加入路径以便导入数据文件
sys.path.insert(0, r"D:\RoboX\02_股权融资")
from data1_0907 import (
    TITLE, OUTPUT_PATH, PARTICIPANTS,
    SUMMARY_1, SUMMARY_2, SUMMARY_3, SUMMARY_4,
    CH1_TITLE, CH1_CONTENT,
    CH2_TITLE, CH2_CONTENT,
    CH3_TITLE, CH3_CONTENT,
    CH4_TITLE, CH4_CONTENT,
    CH5_TITLE, CH5_CONTENT,
)
from data2_0907 import (
    CH6_TITLE, CH6_CONTENT,
    CH7_TITLE, CH7_CONTENT,
    CH8_TITLE, CH8_CONTENT,
    CH9_TITLE, CH9_CONTENT,
    CH10_TITLE, CH10_CONTENT,
    CH11_TITLE, CH11_CONTENT,
    CH12_TITLE, CH12_CONTENT,
)
from data3_0907 import (
    CH13_TITLE, CH13_CONTENT,
    CH14_TITLE, CH14_CONTENT,
    CH15_TITLE, CH15_CONTENT,
    CH16_TITLE, CH16_CONTENT,
    CH17_TITLE, CH17_CONTENT,
    CH18_TITLE, CH18_CONTENT,
)

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ============ 样式定义 ============
TITLE_FONT_SIZE = Pt(16)
HEADING_FONT_SIZE = Pt(13)
BODY_FONT_SIZE = Pt(11)
NOTE_FONT_SIZE = Pt(10)

# 颜色定义
COLOR_TITLE = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝色
COLOR_HEADING = RGBColor(0x1F, 0x4E, 0x79)  # 深蓝色
COLOR_NOTE = RGBColor(0x80, 0x80, 0x80)     # 灰色

# 发言人颜色
SPEAKER_COLORS = {
    "明修": RGBColor(0x70, 0x30, 0xA0),     # 紫色（核心团队/CTO）
    "田明": RGBColor(0x00, 0x80, 0x80),     # 青色（联合创始人）
    "Anson": RGBColor(0x00, 0x00, 0x00),    # 黑色
    "方爱华": RGBColor(0xC0, 0x00, 0x00),   # 深红色（外部嘉宾）
}

FONT_NAME = "微软雅黑"
FONT_NAME_EN = "Microsoft YaHei"


def set_run_font(run, size=BODY_FONT_SIZE, bold=False, color=None, name=FONT_NAME):
    """设置run的字体"""
    run.font.name = name
    run.font.size = size
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # 中文字体设置
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), FONT_NAME_EN)
    rFonts.set(qn('w:hAnsi'), FONT_NAME_EN)


def set_paragraph_format(paragraph, line_spacing=1.5, space_after=Pt(6), space_before=Pt(0)):
    """设置段落格式"""
    pf = paragraph.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_after = space_after
    pf.space_before = space_before


def add_title(doc, text):
    """添加文档标题（居中，16pt，深蓝色加粗）"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(12))
    run = p.add_run(text)
    set_run_font(run, size=TITLE_FONT_SIZE, bold=True, color=COLOR_TITLE)


def add_note(doc, text):
    """添加会议信息（灰色斜体10pt，左缩进1cm）"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(6))
    run = p.add_run(text)
    set_run_font(run, size=NOTE_FONT_SIZE, color=COLOR_NOTE)
    run.italic = True


def add_participants(doc, participants):
    """添加参会人名单（不含职务）"""
    text = "参会人：" + "、".join(participants)
    add_note(doc, text)


def add_summary_heading(doc, text):
    """添加沟通总结分组标题（如'一、路演主线'）"""
    p = doc.add_paragraph()
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(6), space_before=Pt(8))
    run = p.add_run(text)
    set_run_font(run, size=HEADING_FONT_SIZE, bold=True, color=COLOR_HEADING)


def add_summary_item(doc, text):
    """添加沟通总结条目"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)  # 缩进2字符
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(4))
    # 处理粗体标记
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        # 奇数索引为粗体
        set_run_font(run, size=BODY_FONT_SIZE, bold=(i % 2 == 1))


def add_heading(doc, text):
    """添加章节标题（深蓝色加粗，下划线）"""
    p = doc.add_paragraph()
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(8), space_before=Pt(12))
    run = p.add_run(text)
    set_run_font(run, size=HEADING_FONT_SIZE, bold=True, color=COLOR_HEADING)
    # 添加底部边框
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_dialogue(doc, speaker, text):
    """添加对话段落"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(4))
    # 发言人
    speaker_run = p.add_run(f"{speaker}：")
    speaker_color = SPEAKER_COLORS.get(speaker, RGBColor(0x00, 0x00, 0x00))
    set_run_font(speaker_run, size=BODY_FONT_SIZE, bold=True, color=speaker_color)
    # 正文
    text_run = p.add_run(text)
    set_run_font(text_run, size=BODY_FONT_SIZE)


def setup_document(doc):
    """设置文档全局样式"""
    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)
    # 默认样式
    style = doc.styles['Normal']
    style.font.name = FONT_NAME_EN
    style.font.size = BODY_FONT_SIZE


def build_document(output_path):
    """构建Word文档"""
    doc = Document()
    setup_document(doc)

    # ============ 文档标题 ============
    add_title(doc, TITLE)

    # ============ 会议信息 ============
    add_note(doc, "会议时间：2026年9月7日")
    add_note(doc, "会议性质：深创投投委方爱华专场路演")
    add_note(doc, "主讲人：明修（项目创始人/CTO）")
    add_participants(doc, PARTICIPANTS)
    add_note(doc, "主题：基于更新版BP的完整路演 + 投委深度Q&A（约2小时17分）")

    # ============ 沟通总结 ============
    p = doc.add_paragraph()
    set_paragraph_format(p, line_spacing=1.5, space_after=Pt(8), space_before=Pt(12))
    run = p.add_run("沟通总结")
    set_run_font(run, size=HEADING_FONT_SIZE, bold=True, color=COLOR_HEADING)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)

    # 总结1：路演主线
    add_summary_heading(doc, "一、路演主线")
    for item in SUMMARY_1:
        add_summary_item(doc, item)

    # 总结2：技术亮点
    add_summary_heading(doc, "二、技术亮点")
    for item in SUMMARY_2:
        add_summary_item(doc, item)

    # 总结3：Q&A核心议题
    add_summary_heading(doc, "三、Q&A核心议题")
    for item in SUMMARY_3:
        add_summary_item(doc, item)

    # 总结4：后续行动与判断
    add_summary_heading(doc, "四、后续行动与判断")
    for item in SUMMARY_4:
        add_summary_item(doc, item)

    # ============ 转写正文 ============
    chapters = [
        (CH1_TITLE, CH1_CONTENT),
        (CH2_TITLE, CH2_CONTENT),
        (CH3_TITLE, CH3_CONTENT),
        (CH4_TITLE, CH4_CONTENT),
        (CH5_TITLE, CH5_CONTENT),
        (CH6_TITLE, CH6_CONTENT),
        (CH7_TITLE, CH7_CONTENT),
        (CH8_TITLE, CH8_CONTENT),
        (CH9_TITLE, CH9_CONTENT),
        (CH10_TITLE, CH10_CONTENT),
        (CH11_TITLE, CH11_CONTENT),
        (CH12_TITLE, CH12_CONTENT),
        (CH13_TITLE, CH13_CONTENT),
        (CH14_TITLE, CH14_CONTENT),
        (CH15_TITLE, CH15_CONTENT),
        (CH16_TITLE, CH16_CONTENT),
        (CH17_TITLE, CH17_CONTENT),
        (CH18_TITLE, CH18_CONTENT),
    ]

    for title, content in chapters:
        add_heading(doc, title)
        for speaker, text in content:
            add_dialogue(doc, speaker, text)

    # 保存
    doc.save(output_path)
    print(f"[OK] 文档已生成：{output_path}")


if __name__ == "__main__":
    build_document(OUTPUT_PATH)